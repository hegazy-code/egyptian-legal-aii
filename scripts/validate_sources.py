#!/usr/bin/env python3
"""
Validate data/inventory/legal_sources.xlsx against the gap rules already
flagged in that registry (see docs/legal-source-policy.md).

REAL, working implementation -- reads the actual .xlsx via openpyxl and
prints a real report. No mocked/fake output. Flags any row where:
  - "Official Source" is empty, OR contains "NOT YET CONFIRMED" (the
    registry's own convention for a not-yet-Tier-1 source, see
    scripts/_build_initial_registry.py's LAW-003/LAW-004 rows), OR
  - "Publication Date" is empty, OR
  - "Effective Date" is empty, OR
  - "Current Status" == "UNKNOWN"

Usage:
    python scripts/validate_sources.py
    python scripts/validate_sources.py --xlsx data/inventory/legal_sources.xlsx
    python scripts/validate_sources.py --json   # machine-readable report
"""
from __future__ import annotations

import argparse
import json
import sys
from pathlib import Path
from typing import Any, Dict, List, Optional

try:
    import openpyxl
except ImportError:  # pragma: no cover
    print(
        "ERROR: openpyxl is not installed. Run: pip install -r scripts/requirements.txt",
        file=sys.stderr,
    )
    sys.exit(1)

DEFAULT_XLSX_PATH = Path("data/inventory/legal_sources.xlsx")

REQUIRED_COLUMNS = [
    "ID",
    "Legal Area",
    "Law Name",
    "Official Source",
    "Publication Date",
    "Effective Date",
    "Current Status",
]


def _is_blank(value: Any) -> bool:
    return value is None or (isinstance(value, str) and not value.strip())


def load_rows(xlsx_path: Path) -> List[Dict[str, Any]]:
    """Read the registry sheet into a list of {column_name: value} dicts."""
    if not xlsx_path.exists():
        raise FileNotFoundError(
            f"{xlsx_path} does not exist. Nothing to validate -- run "
            f"scripts/_build_initial_registry.py first if this is a fresh checkout."
        )

    wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    ws = wb.active

    header_row = next(ws.iter_rows(min_row=1, max_row=1, values_only=True))
    headers = [str(h) if h is not None else "" for h in header_row]

    missing_cols = [c for c in REQUIRED_COLUMNS if c not in headers]
    if missing_cols:
        raise ValueError(
            f"{xlsx_path} is missing expected column(s): {missing_cols}. "
            f"Found columns: {headers}"
        )

    rows: List[Dict[str, Any]] = []
    for raw_row in ws.iter_rows(min_row=2, values_only=True):
        if raw_row is None or all(_is_blank(v) for v in raw_row):
            continue  # skip fully blank trailing rows
        row = {headers[i]: raw_row[i] for i in range(len(headers)) if i < len(raw_row)}
        rows.append(row)

    return rows


def find_gaps(row: Dict[str, Any]) -> List[str]:
    """Return a list of human-readable gap descriptions for one registry row."""
    gaps: List[str] = []

    official_source = row.get("Official Source")
    if _is_blank(official_source):
        gaps.append("Official Source is empty")
    elif "NOT YET CONFIRMED" in str(official_source).upper():
        gaps.append(
            "Official Source is not yet Tier-1 confirmed "
            f"(currently: {str(official_source)[:80]!r})"
        )

    if _is_blank(row.get("Publication Date")):
        gaps.append("Publication Date is empty")

    if _is_blank(row.get("Effective Date")):
        gaps.append("Effective Date is empty")

    status = row.get("Current Status")
    if status is not None and str(status).strip().upper() == "UNKNOWN":
        gaps.append("Current Status is UNKNOWN")

    return gaps


def build_report(rows: List[Dict[str, Any]]) -> Dict[str, Any]:
    """Build the full validation report as a plain dict (JSON-serializable)."""
    flagged = []
    for row in rows:
        gaps = find_gaps(row)
        if gaps:
            flagged.append(
                {
                    "id": row.get("ID"),
                    "law_name": row.get("Law Name"),
                    "legal_area": row.get("Legal Area"),
                    "gaps": gaps,
                }
            )

    return {
        "total_rows": len(rows),
        "rows_with_gaps": len(flagged),
        "rows_clean": len(rows) - len(flagged),
        "flagged": flagged,
    }


def print_human_report(report: Dict[str, Any]) -> None:
    print(f"Checked {report['total_rows']} row(s) in the legal sources registry.")
    print(f"  {report['rows_clean']} row(s) need no action right now.")
    print(f"  {report['rows_with_gaps']} row(s) still need Tier-1 verification work:\n")

    for entry in report["flagged"]:
        print(f"- [{entry['id']}] {entry['law_name']}  (area: {entry['legal_area']})")
        for gap in entry["gaps"]:
            print(f"    - {gap}")
        print()

    if not report["flagged"]:
        print("No gaps found. (Every row has a confirmed Official Source, both dates, "
              "and a known Current Status.)")


def build_arg_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(
        description=(
            "Report registry rows in data/inventory/legal_sources.xlsx that still "
            "need Tier-1 source verification (empty Official Source, empty "
            "Publication/Effective Date, or Current Status == UNKNOWN)."
        )
    )
    parser.add_argument(
        "--xlsx",
        type=Path,
        default=DEFAULT_XLSX_PATH,
        help=f"Path to the registry file (default: {DEFAULT_XLSX_PATH}).",
    )
    parser.add_argument(
        "--json",
        action="store_true",
        help="Print the report as JSON instead of a human-readable summary.",
    )
    return parser


def main(argv: Optional[list] = None) -> int:
    parser = build_arg_parser()
    args = parser.parse_args(argv)

    try:
        rows = load_rows(args.xlsx)
    except (FileNotFoundError, ValueError) as exc:
        print(f"ERROR: {exc}", file=sys.stderr)
        return 1

    report = build_report(rows)

    if args.json:
        print(json.dumps(report, ensure_ascii=False, indent=2, default=str))
    else:
        print_human_report(report)

    return 0


if __name__ == "__main__":
    sys.exit(main())
