"""
One-off script that built data/inventory/legal_sources.xlsx's initial 4 rows.

Kept in scripts/ (not deleted after running) so the exact provenance of the
seed data is inspectable later -- every value below was checked against a
live web search/fetch on 2026-09-16, not generated from memory. Re-run
manually if the registry needs to be rebuilt from scratch; normally you'd
edit the .xlsx directly or extend this script for new rows.
"""
import openpyxl
from openpyxl.styles import Font, Alignment
from datetime import date

COLUMNS = [
    "ID", "Legal Area", "Law Name", "Law Number", "Year", "Document Type",
    "Issuing Authority", "Official Source", "Source URL", "Document URL",
    "Publication Date", "Effective Date", "Expiration Date", "Current Status",
    "Amended By", "Repealed By", "Related Laws", "Language",
    "Collection Date", "Verification Date", "Verified By", "Notes",
]

TODAY = date(2026, 9, 16).isoformat()

ROWS = [
    {
        "ID": "LAW-001",
        "Legal Area": "Criminal",
        "Law Name": "قانون العقوبات (Penal Code)",
        "Law Number": 58,
        "Year": 1937,
        "Document Type": "Law",
        "Issuing Authority": "Egyptian legislature (parliament)",
        "Official Source": "الوقائع المصرية (Official Gazette), Issue 71",
        "Source URL": "https://www.eastlaws.com/legislation-full-text/ar/egypt/law/05-08-1937/no-58?type=1&id=209",
        "Document URL": "",
        "Publication Date": "1937-08-05",
        "Effective Date": "1937-10-15",
        "Expiration Date": "",
        "Current Status": "AMENDED",
        "Amended By": "قانون 141 لسنة 2024 (seen on one secondary source -- mstsharh.net -- NOT yet confirmed at Gazette/eastlaws level; treat as unverified until cross-checked)",
        "Repealed By": "",
        "Related Laws": "قانون الإجراءات الجنائية (Criminal Procedure Law)",
        "Language": "Arabic",
        "Collection Date": TODAY,
        "Verification Date": TODAY,
        "Verified By": "Claude (session) -- cross-checked eastlaws.com against manshurat.org and lawyeregypt.net for law number/year/name agreement",
        "Notes": "Registry-level metadata only. Article-level text NOT yet extracted -- see README Roadmap step 1. eastlaws.com listing (official Gazette-text database) is the strongest source checked; treat as Tier 1 per docs/legal-source-policy.md.",
    },
    {
        "ID": "LAW-002",
        "Legal Area": "Civil",
        "Law Name": "القانون المدني (Civil Code)",
        "Law Number": 131,
        "Year": 1948,
        "Document Type": "Law",
        "Issuing Authority": "Egyptian legislature (parliament)",
        "Official Source": "الوقائع المصرية (Official Gazette), Issue 108 مكرر",
        "Source URL": "https://www.eastlaws.com/legislation-full-text/ar/egypt/law/29-07-1948/no-131?type=1&id=384",
        "Document URL": "",
        "Publication Date": "1948-07-29",
        "Effective Date": "1949-10-15",
        "Expiration Date": "",
        "Current Status": "ACTIVE",
        "Amended By": "Amended multiple times over decades -- specific amending laws not yet catalogued",
        "Repealed By": "",
        "Related Laws": "قانون المرافعات المدنية والتجارية (Civil and Commercial Procedure Law)",
        "Language": "Arabic",
        "Collection Date": TODAY,
        "Verification Date": TODAY,
        "Verified By": "Claude (session) -- cross-checked eastlaws.com against manshurat.org for law number/year/name agreement",
        "Notes": "Registry-level metadata only. Article-level text NOT yet extracted. eastlaws.com listing checked directly; treat as Tier 1.",
    },
    {
        "ID": "LAW-003",
        "Legal Area": "Companies",
        "Law Name": "قانون شركات المساهمة وشركات التوصية بالأسهم والشركات ذات المسئولية المحدودة (Companies Law)",
        "Law Number": 159,
        "Year": 1981,
        "Document Type": "Law",
        "Issuing Authority": "Egyptian legislature (parliament)",
        "Official Source": "NOT YET CONFIRMED at eastlaws.com/Gazette level -- cross-checked only via secondary publications",
        "Source URL": "https://manshurat.org/node/31504 ; https://draminlotfyoffice.com/details/832",
        "Document URL": "",
        "Publication Date": "",
        "Effective Date": "",
        "Expiration Date": "",
        "Current Status": "UNKNOWN",
        "Amended By": "Known to have been amended multiple times (referenced as 'وفقاً لآخر تعديل' by secondary sources) -- specific amending laws not yet catalogued",
        "Repealed By": "",
        "Related Laws": "قانون الشركات ذات الشخص الواحد",
        "Language": "Arabic",
        "Collection Date": TODAY,
        "Verification Date": TODAY,
        "Verified By": "Claude (session) -- law number/year agreement confirmed across 4 independent secondary sources (manshurat.org, laweg.net, draminlotfyoffice.com, alberonsy.com); NOT yet confirmed against an official Gazette-level source",
        "Notes": "NEEDS a Tier-1 (eastlaws.com or official Gazette) confirmation before being treated as fully verified -- currently only cross-checked secondary sources. Do not cite as authoritative yet. Article-level text NOT extracted.",
    },
    {
        "ID": "LAW-004",
        "Legal Area": "Consumer Protection",
        "Law Name": "قانون حماية المستهلك (Consumer Protection Law)",
        "Law Number": 181,
        "Year": 2018,
        "Document Type": "Law",
        "Issuing Authority": "Egyptian legislature (parliament)",
        "Official Source": "NOT YET CONFIRMED at eastlaws.com/Gazette level -- cross-checked only via secondary publications",
        "Source URL": "https://manshurat.org/node/31479 ; https://lawhub.info/eg/?p=11330",
        "Document URL": "",
        "Publication Date": "",
        "Effective Date": "",
        "Expiration Date": "",
        "Current Status": "UNKNOWN",
        "Amended By": "",
        "Repealed By": "",
        "Related Laws": "اللائحة التنفيذية لقانون حماية المستهلك (Executive Regulations)",
        "Language": "Arabic",
        "Collection Date": TODAY,
        "Verification Date": TODAY,
        "Verified By": "Claude (session) -- law number/year agreement confirmed across 4 independent secondary sources (manshurat.org, laweg.net, lawhub.info, alberonsy.com); NOT yet confirmed against an official Gazette-level source",
        "Notes": "NEEDS a Tier-1 (eastlaws.com or official Gazette) confirmation before being treated as fully verified. Article-level text NOT extracted.",
    },
]


def build():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "legal_sources"

    for col_idx, header in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center")

    for row_idx, row in enumerate(ROWS, start=2):
        for col_idx, header in enumerate(COLUMNS, start=1):
            ws.cell(row=row_idx, column=col_idx, value=row.get(header, ""))

    for col_idx, header in enumerate(COLUMNS, start=1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(col_idx)].width = max(14, min(45, len(header) + 4))

    ws.freeze_panes = "A2"

    out_path = "data/inventory/legal_sources.xlsx"
    wb.save(out_path)
    print(f"Wrote {out_path} with {len(ROWS)} rows")


if __name__ == "__main__":
    build()
